#!/usr/bin/env python3
"""Build hkb/data/taxa.json from the BLP catalogue.

Taxon list only: species, genus, family, order.
Do not add photo counts, folder paths, or image URLs.
Photo hosting is a later discussion.
"""

from __future__ import annotations

import json
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

CATALOGUE = Path(
    "/Users/benjaminlau/Library/CloudStorage/GoogleDrive-lautszonbenjamin@gmail.com"
    "/My Drive/Nature Photo Database/Flora Photo Database (by species)"
    "/BLP_catalogue_20261020.xlsx"
)
OUT = Path(__file__).resolve().parent / "taxa.json"

# Keys that must never appear. Photo inventory is not this file.
FORBIDDEN = ("photo", "n_photos", "n_dates", "folder", "image", "url")


def genus_of(sci: str) -> str:
    # Underscore is a catalogue quirk (one binomial), not a genus character.
    text = str(sci or "").replace("×", " × ").replace("_", " ")
    parts = [p for p in text.split() if p]
    if not parts:
        return ""
    if parts[0] in ("×", "x", "X"):
        return parts[1] if len(parts) > 1 else ""
    return parts[0]


def main() -> None:
    wb = load_workbook(CATALOGUE, read_only=True, data_only=True)

    order_seq = []
    for row in wb["_order_ranges"].iter_rows(min_row=2, values_only=True):
        if row and row[1] == "ORDER" and row[0]:
            order_seq.append(row[0])

    families_by_order = defaultdict(list)
    for row in wb["_family_ranges"].iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        families_by_order[row[1]].append((row[3] or 0, row[0]))
    for order in families_by_order:
        families_by_order[order].sort()

    taxa = []
    for row in wb["Catalogue"].iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        blp = str(row[0]).replace("BLP-", "").zfill(5)
        sci = str(row[3] or "")
        alt = str(row[5] or "").strip()
        item = {
            "blp": blp,
            "status": row[1] or "",
            "rank": row[2] or "",
            "sci": sci,
            "chi": row[4] or "",
            "genus": genus_of(sci),
            "family": row[20] or "",
            "order": row[26] or "",
        }
        if alt:
            item["alt"] = alt
        taxa.append(item)

    taxa.sort(key=lambda t: int(t["blp"]))

    fam_in_ranges = {name for names in families_by_order.values() for _, name in names}
    cat_families = {t["family"] for t in taxa}
    missing = sorted(cat_families - fam_in_ranges)
    extra = sorted(fam_in_ranges - cat_families)
    if missing or extra:
        raise SystemExit(f"family mismatch missing={missing} extra={extra}")

    orders = []
    for name in order_seq:
        fams = [n for _, n in families_by_order.get(name, []) if n in cat_families]
        if not fams:
            raise SystemExit(f"order with no families: {name}")
        orders.append({"name": name, "families": fams})

    cat_orders = {t["order"] for t in taxa}
    if cat_orders != {o["name"] for o in orders}:
        raise SystemExit(f"order mismatch {sorted(cat_orders ^ {o['name'] for o in orders})}")

    for t in taxa:
        if not t["genus"] or not t["family"] or not t["order"] or not t["sci"]:
            raise SystemExit(f"blank rank field {t['blp']}")
        blob = json.dumps(t)
        for bad in FORBIDDEN:
            if bad in blob:
                raise SystemExit(f"forbidden field-like text in {t['blp']}: {bad}")

    genera = {t["genus"] for t in taxa}
    genus_families = defaultdict(set)
    for t in taxa:
        genus_families[t["genus"]].add(t["family"])
    split = {g: sorted(fs) for g, fs in genus_families.items() if len(fs) > 1}

    payload = {
        "source": "BLP_catalogue_20261020.xlsx",
        "built": date.today().isoformat(),
        "note": "Taxon list only. Photo counts and image URLs are intentionally absent.",
        "n_taxa": len(taxa),
        "n_species": sum(1 for t in taxa if t["rank"] == "SPECIES"),
        "n_genera": len(genera),
        "n_families": len(cat_families),
        "n_orders": len(orders),
        "orders": orders,
        "taxa": taxa,
    }
    text = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    for bad in ("n_photos", "photo_folder", "n_dates"):
        if bad in text:
            raise SystemExit(f"{bad} leaked into taxa.json")
    OUT.write_text(text, encoding="utf-8")
    print(f"wrote {OUT}")
    print(
        f"taxa={payload['n_taxa']} species={payload['n_species']} "
        f"genera={payload['n_genera']} families={payload['n_families']} "
        f"orders={payload['n_orders']} bytes={OUT.stat().st_size}"
    )
    print(f"genus_in_two_families={len(split)}")
    for g, fs in sorted(split.items()):
        print(f"  {g}: {fs}")
    samples = [
        "Acer laevigatum",
        "Bauhinia x blakeana",
        "Bombax ceiba",
        "Citrus × limon",
        "Dicliptera sp.",
        "Iris series Hexagonae (red-yellow)",
    ]
    by_sci = {t["sci"]: t for t in taxa}
    for name in samples:
        t = by_sci.get(name)
        print("SAMPLE", name, "MISSING" if not t else f"{t['blp']} {t['genus']} {t['family']} {t['order']} {t['rank']}")


if __name__ == "__main__":
    main()
